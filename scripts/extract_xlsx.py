#!/usr/bin/env python3
"""Extract structure from .xlsx files — sheet names, headers, sample rows."""
import sys, os

def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_xlsx.py <file.xlsx>")
        sys.exit(1)

    from openpyxl import load_workbook

    path = sys.argv[1]
    wb = load_workbook(path, data_only=True, read_only=True)

    print(f"SHEET_COUNT:{len(wb.sheetnames)}")

    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n===SHEET:{name}===")
        print(f"DIMENSIONS:{ws.dimensions}")

        # Print first 20 rows
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 20:
                remaining = ws.max_row - 20 if ws.max_row else "?"
                print(f"... (剩余 {remaining} 行)")
                break
            vals = [str(v).strip() if v is not None else "" for v in row]
            line = " | ".join(vals)
            if line.strip():
                print(line)

    wb.close()

if __name__ == '__main__':
    main()
