#!/usr/bin/env python3
"""Extract text from .pptx (or .ppt via LibreOffice) into structured plain text."""
from pptx import Presentation
import sys, os, subprocess, glob

def convert_ppt_to_pptx(path):
    """Convert .ppt (old binary) to .pptx using LibreOffice headless."""
    outdir = os.path.dirname(path) or "/tmp"
    result = subprocess.run(
        ["soffice", "--headless", "--convert-to", "pptx", "--outdir", outdir, path],
        capture_output=True, text=True, timeout=120
    )
    if result.returncode != 0:
        print(f"ERROR: LibreOffice conversion failed:\n{result.stderr}", file=sys.stderr)
        sys.exit(1)
    basename = os.path.splitext(os.path.basename(path))[0]
    outpath = os.path.join(outdir, basename + ".pptx")
    if not os.path.exists(outpath):
        for f in glob.glob(os.path.join(outdir, "*.pptx")):
            print(f"DEBUG: found {f}", file=sys.stderr)
        sys.exit(1)
    return outpath

def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_pptx.py <file.pptx|file.ppt>")
        sys.exit(1)

    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"ERROR: File not found: {path}")
        sys.exit(1)

    # If .ppt, convert first
    if path.lower().endswith('.ppt') and not path.lower().endswith('.pptx'):
        print(f"Converting .ppt to .pptx: {path}", file=sys.stderr)
        path = convert_ppt_to_pptx(path)

    prs = Presentation(path)

    print(f"SLIDE_COUNT:{len(prs.slides)}")
    print(f"SLIDE_WIDTH:{prs.slide_width}")
    print(f"SLIDE_HEIGHT:{prs.slide_height}")

    for i, slide in enumerate(prs.slides, 1):
        print(f"\n===SLIDE{i}===")

        img_count = 0

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    is_title = shape.is_placeholder and shape.placeholder_format.idx == 0
                    prefix = "TITLE:" if is_title else "TEXT:"
                    print(f"{prefix}{text}")

            if shape.has_table:
                table = shape.table
                rows = len(table.rows)
                cols = len(table.columns)
                print(f"TABLE:{rows}x{cols}")
                header = " | ".join(cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells)
                print(f"TABLE_HEADER:{header}")
                for r in range(1, min(rows, 11)):
                    row_text = " | ".join(cell.text.strip().replace("\n", " ") for cell in table.rows[r].cells)
                    if row_text.strip():
                        print(f"TABLE_ROW:{row_text}")
                if rows > 11:
                    print(f"TABLE_ROW:... (共 {rows - 1} 行数据)")

            if shape.shape_type == 13:  # Picture
                img_count += 1

            if hasattr(shape, 'chart'):
                print("CHART:1")

        if img_count > 0:
            print(f"IMAGE:{img_count}")

if __name__ == '__main__':
    main()
